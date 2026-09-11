import csv
import io
import json
import re
from typing import Any, Dict, List

try:
    from docx import Document as DocxDocument
except Exception:  # pragma: no cover - optional dependency
    DocxDocument = None

try:
    from pypdf import PdfReader
except Exception:  # pragma: no cover - optional dependency
    PdfReader = None

try:
    from openpyxl import load_workbook
except Exception:  # pragma: no cover - optional dependency
    load_workbook = None


def parse_uploaded_records(data: Any, filename: str = "") -> List[Dict[str, Any]]:
    """Parse many common input formats into a list of record dictionaries."""
    raw = _coerce_bytes(data)
    if not raw:
        return []

    name = (filename or "").lower()

    if name.endswith(".jsonl") or name.endswith(".ndjson"):
        return _parse_json_lines(raw)

    if name.endswith(".json"):
        return _parse_json_payload(raw)

    if name.endswith(".csv"):
        return _parse_delimited(raw, ",")

    if name.endswith(".tsv"):
        return _parse_delimited(raw, "\t")

    if name.endswith(".xlsx") or name.endswith(".xls"):
        return _parse_excel(raw)

    if name.endswith(".docx"):
        return _parse_docx(raw)

    if name.endswith(".pdf"):
        return _parse_pdf(raw)

    if name.endswith(".txt") or name.endswith(".log") or name.endswith(".text"):
        return _parse_text(raw)

    return _parse_by_content(raw)


def _coerce_bytes(data: Any) -> bytes:
    if isinstance(data, bytes):
        return data
    if isinstance(data, bytearray):
        return bytes(data)
    if isinstance(data, str):
        return data.encode("utf-8-sig")
    return b""


def _decode_text(data: bytes) -> str:
    return data.decode("utf-8-sig", errors="ignore")


def _normalize_record(entry: Any) -> Dict[str, Any]:
    if isinstance(entry, dict):
        normalized = {}
        for key, value in entry.items():
            normalized[str(key)] = _normalize_value(value)
        return normalized

    if isinstance(entry, (list, tuple)):
        return {
            "value": _normalize_value(entry),
        }

    return {"content": _normalize_value(entry)}


def _normalize_value(value: Any) -> Any:
    if isinstance(value, (dict, list)):
        return json.loads(json.dumps(value, ensure_ascii=False))
    return value


def _parse_by_content(raw: bytes) -> List[Dict[str, Any]]:
    text = _decode_text(raw).strip()
    if not text:
        return []

    if text.startswith("{") or text.startswith("["):
        return _parse_json_payload(raw)

    if _looks_like_delimited(text):
        delimiter = _detect_delimiter(text)
        return _parse_delimited(raw, delimiter)

    return _parse_text(raw)


def _parse_json_payload(raw: bytes) -> List[Dict[str, Any]]:
    payload = json.loads(_decode_text(raw))

    if isinstance(payload, dict):
        for key, value in payload.items():
            if isinstance(value, list) and value and all(isinstance(item, dict) for item in value):
                return [_normalize_record(item) for item in value]
        return [_normalize_record(payload)]

    if isinstance(payload, list):
        return [_normalize_record(item) for item in payload]

    return [{"content": _normalize_value(payload)}]


def _parse_json_lines(raw: bytes) -> List[Dict[str, Any]]:
    records = []
    for line in _decode_text(raw).splitlines():
        line = line.strip()
        if not line:
            continue
        records.append(_normalize_record(json.loads(line)))
    return records


def _parse_excel(raw: bytes) -> List[Dict[str, Any]]:
    if load_workbook is None:
        raise RuntimeError("openpyxl is required to read Excel files. Install it with pip install openpyxl.")

    workbook = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
    records: List[Dict[str, Any]] = []

    for sheet in workbook.worksheets:
        rows = list(sheet.iter_rows(values_only=True))
        if not rows:
            continue
        headers = [str(cell).strip() if cell is not None else f"column_{idx + 1}" for idx, cell in enumerate(rows[0])]

        for row in rows[1:]:
            row_values = list(row)
            if not any(value not in (None, "", " ") for value in row_values):
                continue
            record = {}
            for idx, value in enumerate(row_values):
                key = headers[idx] if idx < len(headers) else f"column_{idx + 1}"
                record[key] = _normalize_value(value)
            records.append(record)

    return records


def _parse_docx(raw: bytes) -> List[Dict[str, Any]]:
    if DocxDocument is None:
        raise RuntimeError("python-docx is required to read Word files. Install it with pip install python-docx.")

    doc = DocxDocument(io.BytesIO(raw))
    records: List[Dict[str, Any]] = []

    for table in doc.tables:
        if not table.rows:
            continue
        headers = [cell.text.strip() or f"column_{idx + 1}" for idx, cell in enumerate(table.rows[0].cells)]
        for row in table.rows[1:]:
            values = [cell.text.strip() for cell in row.cells]
            if not any(values):
                continue
            record = {}
            for idx, value in enumerate(values):
                record[headers[idx] if idx < len(headers) else f"column_{idx + 1}"] = _normalize_value(value)
            records.append(record)

    if records:
        return records

    for paragraph in doc.paragraphs:
        text = paragraph.text.strip()
        if text:
            records.append({"content": text})

    return records


def _parse_pdf(raw: bytes) -> List[Dict[str, Any]]:
    if PdfReader is None:
        raise RuntimeError("pypdf is required to read PDF files. Install it with pip install pypdf.")

    reader = PdfReader(io.BytesIO(raw))
    text_parts = []
    for page in reader.pages:
        extracted = page.extract_text() or ""
        text_parts.append(extracted)

    return _parse_text("\n".join(text_parts).encode("utf-8"))


def _parse_delimited(raw: bytes, delimiter: str) -> List[Dict[str, Any]]:
    text = _decode_text(raw)
    if not text.strip():
        return []

    reader = csv.DictReader(io.StringIO(text), delimiter=delimiter)
    records = []
    for row in reader:
        if row is None:
            continue
        normalized = {str(key): _normalize_value(value) for key, value in row.items() if key is not None}
        if normalized:
            records.append(normalized)

    if records:
        return records

    # Fallback: treat each non-empty line as a single record
    return [{"content": line.strip()} for line in text.splitlines() if line.strip()]


def _parse_text(raw: bytes) -> List[Dict[str, Any]]:
    text = _decode_text(raw).strip()
    if not text:
        return []

    if _looks_like_delimited(text):
        delimiter = _detect_delimiter(text)
        return _parse_delimited(raw, delimiter)

    paragraphs = re.split(r"\n\s*\n+", text)
    if len(paragraphs) > 1:
        return [{"content": paragraph.strip()} for paragraph in paragraphs if paragraph.strip()]

    # If a plain text file contains a line-based list, preserve each line as a record.
    lines = [line.strip() for line in text.splitlines() if line.strip()]
    if len(lines) > 1:
        return [{"content": line} for line in lines]

    return [{"content": text}]


def _detect_delimiter(text: str) -> str:
    try:
        dialect = csv.Sniffer().sniff(text[:4096], delimiters=",\t;|")
        return dialect.delimiter
    except Exception:
        if "\t" in text:
            return "\t"
        if ";" in text:
            return ";"
        if "|" in text:
            return "|"
        return ","


def _looks_like_delimited(text: str) -> bool:
    if not text or text.startswith("{") or text.startswith("["):
        return False

    lines = [line.strip() for line in text.splitlines() if line.strip()]
    if len(lines) < 2:
        return False

    first = lines[0]
    if "," in first or "\t" in first or ";" in first or "|" in first:
        return True

    return False
